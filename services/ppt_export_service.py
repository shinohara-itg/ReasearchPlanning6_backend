from __future__ import annotations

from io import BytesIO
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from app.schemas.ppt_export import PptShapeItem


def _iter_shapes_recursive(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes_recursive(shape.shapes)


def _apply_text_style(shape, font_name: str = "Arial", font_size: int = 12) -> None:
    if not getattr(shape, "has_text_frame", False):
        return

    try:
        text_frame = shape.text_frame
        text_frame.word_wrap = True

        for paragraph in text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
            for run in paragraph.runs:
                run.font.name = font_name
                run.font.size = Pt(font_size)
                run.font.color.rgb = RGBColor(0, 0, 0)
    except Exception:
        pass


def _set_text_to_named_shape(slide, shape_name: str, text: str) -> bool:
    for shape in _iter_shapes_recursive(slide.shapes):
        if getattr(shape, "name", "") != shape_name:
            continue

        if getattr(shape, "has_text_frame", False):
            shape.text = text
            _apply_text_style(shape)
            return True

        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            try:
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell.text = text
                        _apply_text_style(cell)
                return True
            except Exception:
                return False

        return False

    return False


def export_ppt_from_template_bytes(
    template_bytes: bytes,
    items: Iterable[PptShapeItem],
) -> tuple[BytesIO, dict]:
    prs = Presentation(BytesIO(template_bytes))

    report = {
        "total_items": 0,
        "applied": 0,
        "skipped_empty": 0,
        "slide_oob": 0,
        "shape_not_found": 0,
        "errors": [],
    }

    for item in items:
        report["total_items"] += 1

        try:
            text = item.text or ""
            if not text.strip():
                report["skipped_empty"] += 1
                continue

            if item.slide_index < 0 or item.slide_index >= len(prs.slides):
                report["slide_oob"] += 1
                continue

            slide = prs.slides[item.slide_index]
            ok = _set_text_to_named_shape(
                slide=slide,
                shape_name=item.shape_name,
                text=text,
            )

            if ok:
                report["applied"] += 1
            else:
                report["shape_not_found"] += 1

        except Exception as e:
            report["errors"].append(
                {
                    "slide_index": item.slide_index,
                    "shape_name": item.shape_name,
                    "error": str(e),
                }
            )

    output = BytesIO()
    prs.save(output)
    output.seek(0)

    return output, report